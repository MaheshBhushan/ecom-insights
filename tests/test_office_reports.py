"""Tests for deterministic Excel and PowerPoint report generation."""

from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

from src.reporting.office_reports import (
    DEFAULT_PPTX,
    DEFAULT_XLSX,
    build_excel_report,
    build_powerpoint_report,
    load_report_data,
    validate_outputs,
)

PROJECT_ROOT = Path(__file__).resolve().parents[1]


def test_report_inputs_reconcile_net_category_total():
    data = load_report_data(
        PROJECT_ROOT / "data" / "processed", PROJECT_ROOT / "reports"
    )
    fact = data.frames["fact_sales"]
    included = ~fact["order_status"].isin(
        data.kpis["order_statuses_excluded_from_revenue"]
    )
    assert (
        round(float(fact.loc[included, "net_revenue"].sum()), 2)
        == data.kpis["net_revenue_total"]
    )


def test_generates_structurally_valid_office_reports(tmp_path):
    data = load_report_data(
        PROJECT_ROOT / "data" / "processed", PROJECT_ROOT / "reports"
    )
    xlsx = build_excel_report(data, tmp_path / DEFAULT_XLSX)
    pptx = build_powerpoint_report(data, tmp_path / DEFAULT_PPTX)
    validate_outputs(xlsx, pptx, data)

    workbook = load_workbook(xlsx, data_only=False)
    assert workbook.sheetnames == [
        "Executive Summary",
        "Sales Trends",
        "Categories",
        "Products",
        "Customer Segments",
        "Cohort Retention",
        "Data Dictionary",
    ]
    assert workbook["Executive Summary"]["A5"].value == data.kpis["net_revenue_total"]
    assert len(workbook["Sales Trends"]._charts) == 1
    assert len(workbook["Categories"]._charts) == 1

    presentation = Presentation(pptx)
    assert len(presentation.slides) == 7
    visible_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text_frame") and shape.has_text_frame
    )
    assert "Revenue is healthy, but one in four orders is cancelled" in visible_text
    assert f"${data.kpis['net_revenue_total']:,.0f}" in visible_text
    assert "Customer retention is a data-quality warning, not a win" in visible_text
